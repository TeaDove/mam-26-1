import sys
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from pydantic import BaseModel, TypeAdapter

from lab2.dto.models import (
    ChunkingMetrics,
    CleaningMetrics,
    GoldPageMetrics,
    GraphComparison,
    GraphReport,
    GraphSummary,
    NormalizationMetrics,
    RetrievalMetrics,
    TokenizationMetrics,
    TraversalSummary,
    VectorizationMetrics,
)

ROOT = Path(__file__).resolve().parents[1]
METRICS = ROOT / "results" / "metrics"
FIGURES = ROOT / "results" / "figures"
AUTHORS = ROOT / "authors.txt"

INK = RGBColor(0x1F, 0x23, 0x28)
MUTED = RGBColor(0x5F, 0x5E, 0x5A)
CLEAN = RGBColor(0x2A, 0x78, 0xD6)
DIRTY = RGBColor(0xEB, 0x68, 0x34)
TINT = RGBColor(0xF1, 0xF4, 0xF8)
LINE = RGBColor(0xD9, 0xD8, 0xD3)
CODE_BG = RGBColor(0xF4, 0xF4, 0xF2)
MARK_COLOR = RGBColor(0x00, 0x83, 0x00)
MARKS = {"L": "\u22c6", "D": "\u22c4", "H": "\u22b9"}
FONT = "Calibri"


def load[T: BaseModel](name: str, model: type[T]) -> T:
    return model.model_validate_json((METRICS / name).read_text(encoding="utf-8"))


def load_columns[T: BaseModel](name: str, model: type[T]) -> dict[str, T]:
    return TypeAdapter(dict[str, model]).validate_json((METRICS / name).read_text(encoding="utf-8"))


def bfs_triple(summary: TraversalSummary) -> str:
    values = (summary.bfs1_mean_reached, summary.bfs2_mean_reached, summary.bfs3_mean_reached)
    return " / ".join(num(value, 0) for value in values)


def pct(value: float, digits: int = 0) -> str:
    return f"{value * 100:.{digits}f}%".replace(".", ",")


def num(value: float, digits: int = 2) -> str:
    return f"{value:.{digits}f}".replace(".", ",")


def split_mark(text: str) -> tuple[str, str | None]:
    if len(text) > 2 and text[-2] == "@" and text[-1] in MARKS:
        return text[:-2], text[-1]
    return text, None


def add_mark(paragraph: object, kind: str | None, size: float) -> None:
    if kind is None:
        return
    run = paragraph.add_run()
    run.text = " " + MARKS[kind]
    run.font.size = Pt(max(11, size))
    run.font.bold = False
    run.font.color.rgb = MARK_COLOR


class Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.blank = self.prs.slide_layouts[6]

    def slide(self, title: str | None, notes: str) -> object:
        slide = self.prs.slides.add_slide(self.blank)
        if title:
            self.text(slide, title, 0.6, 0.45, 12.1, 0.8, size=28, bold=True)
        slide.notes_slide.notes_text_frame.text = notes
        return slide

    def text(
        self,
        slide: object,
        text: str | list[str],
        left: float,
        top: float,
        width: float,
        height: float,
        size: int = 16,
        bold: bool = False,
        color: RGBColor = INK,
        align: PP_ALIGN = PP_ALIGN.LEFT,
        font: str = FONT,
        spacing: int = 6,
    ) -> object:
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        lines = text if isinstance(text, list) else [text]
        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = align
            paragraph.space_after = Pt(spacing)
            line, kind = split_mark(line)
            run = paragraph.add_run()
            run.text = line
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
            run.font.name = font
            add_mark(paragraph, kind, size)
        return box

    def card(self, slide: object, left: float, top: float, width: float, height: float) -> None:
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
        )
        shape.adjustments[0] = 0.06
        shape.fill.solid()
        shape.fill.fore_color.rgb = TINT
        shape.line.fill.background()
        shape.shadow.inherit = False

    def stat(self, slide: object, left: float, top: float, before: str, after: str, label: str) -> None:
        self.card(slide, left, top, 3.8, 1.75)
        box = slide.shapes.add_textbox(Inches(left + 0.3), Inches(top + 0.22), Inches(3.3), Inches(0.75))
        frame = box.text_frame
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        paragraph = frame.paragraphs[0]
        for value, color, size in ((before, DIRTY, 26), ("  →  ", MUTED, 22), (after, CLEAN, 34)):
            run = paragraph.add_run()
            run.text = value
            run.font.size = Pt(size)
            run.font.bold = True
            run.font.color.rgb = color
            run.font.name = FONT
        self.text(slide, label, left + 0.3, top + 1.0, 3.45, 0.7, size=13, color=MUTED)

    def table(
        self,
        slide: object,
        rows: list[list[str]],
        left: float,
        top: float,
        widths: list[float],
        size: int = 13,
        value_columns: tuple[int, int] = (-2, -1),
        row_height: float = 0.38,
    ) -> None:
        shape = slide.shapes.add_table(
            len(rows), len(widths), Inches(left), Inches(top), Inches(sum(widths)), Inches(row_height * len(rows))
        )
        table = shape.table
        dirty_column, clean_column = (column % len(widths) for column in value_columns)
        for column, width in enumerate(widths):
            table.columns[column].width = Inches(width)
        for r, row in enumerate(rows):
            table.rows[r].height = Inches(row_height)
            for c, value in enumerate(row):
                cell = table.cell(r, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = TINT if r == 0 else RGBColor(0xFF, 0xFF, 0xFF)
                cell.margin_left = cell.margin_right = Inches(0.08)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                frame = cell.text_frame
                frame.paragraphs[0].text = ""
                value, kind = split_mark(value) if c == 0 and r > 0 else (value, None)
                run = frame.paragraphs[0].add_run()
                run.text = value
                run.font.size = Pt(size)
                run.font.name = FONT
                run.font.bold = r == 0
                if r > 0 and c == clean_column:
                    run.font.color.rgb = CLEAN
                elif r > 0 and c == dirty_column:
                    run.font.color.rgb = DIRTY
                else:
                    run.font.color.rgb = INK if r > 0 else MUTED
                numeric = c in (dirty_column, clean_column)
                frame.paragraphs[0].alignment = PP_ALIGN.RIGHT if numeric else PP_ALIGN.LEFT
                add_mark(frame.paragraphs[0], kind, size)

    def bars(
        self,
        slide: object,
        title: str,
        categories: list[str],
        dirty: list[float],
        clean: list[float],
        left: float,
        top: float,
        width: float,
        height: float,
        number_format: str = "0",
    ) -> None:
        data = CategoryChartData()
        data.categories = categories
        data.add_series("грязный", dirty)
        data.add_series("чистый", clean)
        frame = slide.shapes.add_chart(
            XL_CHART_TYPE.BAR_CLUSTERED, Inches(left), Inches(top), Inches(width), Inches(height), data
        )
        chart = frame.chart
        chart.has_title = True
        title, kind = split_mark(title)
        chart.chart_title.text_frame.text = title
        title_run = chart.chart_title.text_frame.paragraphs[0].runs[0]
        title_run.font.size = Pt(14)
        title_run.font.bold = True
        title_run.font.color.rgb = INK
        add_mark(chart.chart_title.text_frame.paragraphs[0], kind, 14)
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(12)
        chart.legend.font.color.rgb = MUTED
        plot = chart.plots[0]
        plot.gap_width = 60
        plot.overlap = -10
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.number_format = number_format
        labels.number_format_is_linked = False
        labels.position = XL_LABEL_POSITION.OUTSIDE_END
        labels.font.size = Pt(11)
        labels.font.color.rgb = INK
        for series, color in zip(plot.series, (DIRTY, CLEAN), strict=True):
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = color
        category_axis = chart.category_axis
        category_axis.tick_labels.font.size = Pt(12)
        category_axis.tick_labels.font.color.rgb = INK
        category_axis.format.line.color.rgb = LINE
        category_axis.reverse_order = True
        value_axis = chart.value_axis
        value_axis.visible = False
        value_axis.has_major_gridlines = False

    def picture(self, slide: object, path: Path, left: float, top: float, height: float) -> None:
        slide.shapes.add_picture(str(path), Inches(left), Inches(top), height=Inches(height))

    def arrow_row(self, slide: object, label: str, steps: list[str], top: float, color: RGBColor) -> None:
        self.text(slide, label, 0.6, top + 0.2, 1.4, 0.4, size=16, bold=True, color=color)
        left, width, gap, height = 2.1, 1.28, 0.26, 0.72
        for index, step in enumerate(steps):
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
            )
            shape.adjustments[0] = 0.15
            shape.fill.solid()
            shape.fill.fore_color.rgb = TINT
            shape.line.color.rgb = color
            shape.line.width = Pt(1.25)
            shape.shadow.inherit = False
            frame = shape.text_frame
            frame.margin_left = frame.margin_right = Inches(0.02)
            frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            frame.word_wrap = False
            paragraph = frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.CENTER
            step, kind = split_mark(step)
            run = paragraph.add_run()
            run.text = step
            run.font.size = Pt(10)
            run.font.color.rgb = INK
            run.font.name = FONT
            add_mark(paragraph, kind, 10)
            if index < len(steps) - 1:
                arrow = self.text(
                    slide, "→", left + width, top, gap, height, size=14, color=MUTED, align=PP_ALIGN.CENTER
                )
                arrow.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            left += width + gap

    def code(self, slide: object, lines: list[str], top: float) -> None:
        height = 0.32 * len(lines) + 0.3
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(top), Inches(12.1), Inches(height)
        )
        shape.adjustments[0] = 0.08
        shape.fill.solid()
        shape.fill.fore_color.rgb = CODE_BG
        shape.line.fill.background()
        shape.shadow.inherit = False
        self.text(slide, lines, 0.85, top + 0.15, 11.7, height - 0.2, size=12, color=INK, font="Consolas", spacing=2)

    def legend(self, slide: object, left: float, top: float) -> None:
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(12.1), Inches(0.4))
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.LEFT
        run = paragraph.add_run()
        run.text = "Метки у метрик:"
        run.font.size = Pt(14)
        run.font.color.rgb = MUTED
        run.font.name = FONT
        for kind, meaning in (("L", "LLM"), ("D", "алгоритм"), ("H", "гибрид: обход графа + LLM")):
            add_mark(paragraph, kind, 18)
            run = paragraph.add_run()
            run.text = f" — {meaning}   "
            run.font.size = Pt(14)
            run.font.color.rgb = MUTED
            run.font.name = FONT

    def background(self, slide: object, path: Path) -> None:
        picture = slide.shapes.add_picture(str(path), 0, 0, width=self.prs.slide_width, height=self.prs.slide_height)
        tree = slide.shapes._spTree
        tree.remove(picture._element)
        tree.insert(2, picture._element)

    def save(self, path: Path) -> None:
        total = len(self.prs.slides)
        for number, slide in enumerate(self.prs.slides, start=1):
            self.text(slide, f"{number} / {total}", 11.73, 7.14, 1.0, 0.26, size=11, color=MUTED, align=PP_ALIGN.RIGHT)
        self.prs.save(str(path))


@dataclass
class Metrics:
    cleaning: dict[str, CleaningMetrics]
    normalization: dict[str, NormalizationMetrics]
    chunking: dict[str, ChunkingMetrics]
    tokenization: dict[str, TokenizationMetrics]
    vectors: dict[str, VectorizationMetrics]
    retrieval: dict[str, RetrievalMetrics]
    gold: GoldPageMetrics
    graphs: dict[str, GraphSummary]
    comparison: GraphComparison

    @classmethod
    def load_all(cls) -> Metrics:
        return cls(
            cleaning=load_columns("01_cleaning.json", CleaningMetrics),
            normalization=load_columns("02_normalization.json", NormalizationMetrics),
            chunking=load_columns("03_chunking.json", ChunkingMetrics),
            tokenization=load_columns("04_tokenization.json", TokenizationMetrics),
            vectors=load_columns("05_vectorization.json", VectorizationMetrics),
            retrieval=load_columns("05_retrieval.json", RetrievalMetrics),
            gold=load_columns("06_gold.json", GoldPageMetrics)["total"],
            graphs=load_columns("07_graphs.json", GraphSummary),
            comparison=load("08_graph_comparison.json", GraphComparison),
        )

    @property
    def dirty(self) -> GraphReport:
        return self.comparison.dirty

    @property
    def clean(self) -> GraphReport:
        return self.comparison.clean


def title_slide(deck: Deck, m: Metrics) -> None:
    slide = deck.slide(
        None,
        "Тема: как предобработка текста меняет граф знаний. Корпус — два учебника по контролируемой прокатке стали: "
        "английский обзор Tanaka (1981) и русская статья. Оба текста строят один общий граф. Сравниваем граф из сырого "
        "текста MinerU и граф после предобработки. Это вторая версия: этапы переставлены, качество этапов меряется на "
        "эталоне, все критерии считаются обходами графа.",
    )
    deck.text(slide, "Предобработка текста для графа знаний", 0.8, 2.0, 11.8, 1.0, size=36, bold=True)
    deck.text(
        slide,
        "Грязный и чистый граф по двум учебникам о контролируемой прокатке стали",
        0.8,
        3.25,
        11.8,
        0.6,
        size=22,
        color=MUTED,
    )
    deck.text(
        slide,
        "tanaka1981 (EN, 28 с.) + stat3 (RU, 10 с.) · MinerU → GraphRAG · gpt-4o-mini, bge-m3 · судья gpt-4.1",
        0.8,
        4.1,
        11.8,
        0.5,
        size=16,
        color=MUTED,
    )
    if AUTHORS.exists():
        authors = [line.strip() for line in AUTHORS.read_text(encoding="utf-8").splitlines() if line.strip()]
        deck.text(slide, authors, 7.7, 5.5, 5.0, 1.5, size=14, color=MUTED, align=PP_ALIGN.RIGHT, spacing=2)


def experiment_slide(deck: Deck, m: Metrics) -> None:
    slide = deck.slide(
        "Эксперимент: отличие только в предобработке",
        "Оба графа строит GraphRAG с одинаковыми настройками: gpt-4o-mini, bge-m3, те же промпты и типы сущностей. "
        "Грязный граф — из markdown MinerU, из которого убран только список литературы, иначе сравнение нечестное. "
        "Чистый — после пяти этапов. Очистка и нормализация работают с книгой целиком, chunking режет уже чистый "
        "текст. Качество этапов меряем на четырёх страницах, которые автор выверил вручную по PDF; проверка эталона "
        "командой ещё не завершена, поэтому цифры на эталоне предварительные. Критерии сравнения графов "
        "считаем обходами графа; LLM — только для непротиворечивости.",
    )
    deck.arrow_row(slide, "Грязный", ["MinerU", "− литература", "GraphRAG@L"], 1.8, DIRTY)
    deck.arrow_row(
        slide,
        "Чистый",
        ["MinerU", "Очистка@D", "Нормализация@D", "Chunking@D", "Токенизация@D", "Векторизация@D", "GraphRAG@L"],
        3.0,
        CLEAN,
    )
    dirty, clean = m.graphs["dirty"], m.graphs["clean"]
    deck.text(
        slide,
        [
            "Одинаковые модели, промпты и типы сущностей; оба учебника → один граф.",
            "Эталон: 4 страницы (по 2 из книги), выверенные автором по PDF, проверка командой не завершена — "
            "CER/WER, точность и полнота удалений, точность приведения величин и терминов.",
            f"Графы: грязный {dirty.vertices} вершин / {dirty.edges} рёбер, "
            f"чистый {clean.vertices} / {clean.edges}; у каждой вершины вектор bge-m3.",
        ],
        0.6,
        4.4,
        12.1,
        1.8,
        size=16,
        color=MUTED,
    )
    deck.legend(slide, 0.6, 6.3)


def cleaning_slide(deck: Deck, m: Metrics) -> None:
    cl_d, cl_c, gold = m.cleaning["dirty:total"], m.cleaning["clean:total"], m.gold
    slide = deck.slide(
        "1. Очистка: на эталоне ошибок почти не осталось",
        "Удаляем колонтитулы, номера страниц, ссылки на картинки, HTML-сноски, маркеры цитат и шапку с авторами, "
        f"склеиваем абзацы, разорванные страницей, и переносы. Качество — на эталонных страницах: CER упал с "
        f"{pct(gold.raw_cer, 1)} до {pct(gold.clean_cer, 1)}, удалено {gold.blocks_deleted_correctly} из "
        f"{gold.blocks_should_delete} блоков, которые должны были уйти, и ни одного лишнего. Это хорошо: остаток "
        "CER — OCR-ошибки в индексах таблицы, которые очистка не исправляет. Числа, формулы и ячейки таблиц "
        f"сохранены на {pct(cl_c.numbers_preserved_share)}.",
    )
    deck.stat(slide, 0.6, 1.55, pct(gold.raw_cer, 1), pct(gold.clean_cer, 1), "CER относительно эталона@D")
    deck.stat(slide, 4.75, 1.55, pct(gold.raw_wer, 1), pct(gold.clean_wer, 1), "WER относительно эталона@D")
    deck.stat(
        slide,
        8.9,
        1.55,
        "—",
        f"{gold.blocks_deleted_correctly}/{gold.blocks_should_delete}",
        f"удалено нужных блоков, лишних {gold.blocks_deleted - gold.blocks_deleted_correctly}; "
        f"F1 {pct(gold.deletion_f1)}@D",
    )
    deck.table(
        slide,
        [
            ["Метрика", "сырой", "после очистки"],
            ["Доля шумовых символов@D", pct(cl_d.noise_share, 1), pct(cl_c.noise_share, 2)],
            [
                "Служебные метки / разорванные абзацы@D",
                f"{cl_d.service_marks} / {cl_d.broken_paragraphs}",
                f"{cl_c.service_marks} / {cl_c.broken_paragraphs}",
            ],
            [
                "Переносы в словах / PII@D",
                f"{cl_d.hyphenated_breaks} / {cl_d.pii_hits}",
                f"{cl_c.hyphenated_breaks} / {cl_c.pii_hits}",
            ],
            [
                "Сохранено чисел / формул / ячеек таблиц@D",
                "—",
                f"{pct(cl_c.numbers_preserved_share)} / {pct(cl_c.formulas_preserved_share)} / "
                f"{pct(cl_c.table_cells_preserved_share)}",
            ],
        ],
        0.6,
        3.65,
        [6.6, 2.4, 3.1],
        size=14,
    )
    deck.code(
        slide,
        [
            r're.sub(r"</?(?:strong|u|b|i|em)>", "", text);   re.sub(r"\s*\$\^\{[\d,\s–-]+\}\$", "", text)',
            r're.sub(r"([^\W\d_]+)-\s+([а-яё]+)", r"\1\2", text);   блоки = колонтитулы ∪ r"^\d{1,3}$" → удалить',
        ],
        5.85,
    )


def normalization_slide(deck: Deck, m: Metrics) -> None:
    raw, norm, gold = m.normalization["raw:total"], m.normalization["normalized:total"], m.gold
    slide = deck.slide(
        "2. Нормализация: величины, термины, леммы",
        "Приводим величины к одному виду, расшифровываем аббревиатуры и ставим перед русским термином канонический "
        "английский — заново каждые 3000 символов, поэтому пометка не зависит от будущего разбиения. Английский в "
        "скобках после русского модель игнорирует, а английский первым — нет: так у двух книг появилось "
        f"{norm.bridging_terms} общих терминов. Термины лемматизируются: существительные и пары «прилагательное + "
        "существительное» подряд. Цена — русский текст становится неграмматичным («со скорость»): для графа это "
        "нужно, чтобы формы слова склеились в одну вершину, но читать такой текст хуже. На эталоне все "
        f"{gold.units_expected} величин и {gold.terms_expected} терминов приведены верно — хорошо, но выборка "
        "маленькая.",
    )
    deck.stat(
        slide, 0.6, 1.55, pct(raw.unit_canonical_share), pct(norm.unit_canonical_share), "величин в едином виде@D"
    )
    deck.stat(slide, 0.6, 3.5, str(raw.bridging_terms), str(norm.bridging_terms), "терминов глоссария в обеих книгах@D")
    deck.card(slide, 4.75, 1.55, 7.95, 3.7)
    deck.text(slide, "stat3, было", 5.05, 1.7, 7.3, 0.4, size=14, bold=True, color=DIRTY)
    deck.text(
        slide,
        "…эффект КП, как способ измельчения зерна аустенита путем рекристаллизации… "
        "ускоренное охлаждение после деформации со скоростями 10–30<sup>0</sup>С/сек",
        5.05,
        2.1,
        7.4,
        1.0,
        size=14,
    )
    deck.text(slide, "стало (леммы терминов — ценой грамматики)", 5.05, 3.1, 7.3, 0.4, size=14, bold=True, color=CLEAN)
    deck.text(
        slide,
        "…эффект КП, как способ grain refinement (измельчения зерно) аустенит путем рекристаллизация… "
        "ускоренное охлаждение после deformation (деформация) со скорость 10–30 °C/с",
        5.05,
        3.5,
        7.4,
        1.6,
        size=14,
    )
    deck.text(
        slide,
        f"Эталон: величины {gold.units_correct}/{gold.units_expected}, термины "
        f"{gold.terms_correct}/{gold.terms_expected}; лемматизировано {pct(norm.lemmatized_share or 0, 1)} "
        f"слов (полная лемматизация — {pct(norm.full_lemmatization_share or 0)})@D",
        0.6,
        5.45,
        12.1,
        0.5,
        size=14,
        color=MUTED,
    )
    deck.code(
        slide,
        [
            r're.sub(r"<sup>\s*[0o°]\s*</sup>\s*[СC]", " °C", text)',
            r'periodic(r"\bконтролируем\w* прокатк\w*", lambda m: f"controlled rolling ({m[0]})", every=3000)'
            "   ← glossary.yaml",
        ],
        6.0,
    )


def chunking_slide(deck: Deck, m: Metrics) -> None:
    ch_d, ch_c = m.chunking["dirty_graphrag:total"], m.chunking["structural:total"]
    cutting = round(ch_c.boundaries_cutting_sentence_share * ch_c.boundaries)
    slide = deck.slide(
        "3. Chunking: по структуре уже чистого текста",
        "GraphRAG по умолчанию режет текст окнами по 1200 токенов: каждая граница рвёт предложение, таблица "
        "разрезана. Структурный разбивщик идёт по разделам и абзацам чистого текста, не режет формулы и таблицы и "
        f"ставит границу на конце предложения. Внутри предложения осталось {cutting} из {ch_c.boundaries} границ — "
        "длинный раздел без удобного места. Чанки мельче, поэтому GraphRAG их не перерезает; они сохранены для RAG. "
        "Обратная сторона: из мелких чанков модель извлекает больше вершин на токен — это видно дальше на графах.",
    )
    deck.stat(
        slide,
        0.6,
        1.55,
        pct(ch_d.boundaries_cutting_sentence_share),
        pct(ch_c.boundaries_cutting_sentence_share),
        "границ чанков режут предложение@D",
    )
    deck.stat(
        slide,
        0.6,
        3.5,
        str(ch_d.tables_broken + ch_d.formulas_broken),
        str(ch_c.tables_broken + ch_c.formulas_broken),
        "формул и таблиц разрезано границей@D",
    )
    books = ("tanaka1981", "stat3", "total")
    deck.bars(
        slide,
        "Размер чанка, токенов o200k (среднее)@D",
        ["tanaka1981", "stat3", "всего"],
        [m.chunking[f"dirty_graphrag:{book}"].tokens.mean for book in books],
        [m.chunking[f"structural:{book}"].tokens.mean for book in books],
        4.9,
        1.45,
        7.8,
        3.9,
    )
    deck.text(
        slide,
        f"Чанков: {ch_d.n_chunks} → {ch_c.n_chunks}; максимум {ch_c.tokens.max} токенов.",
        4.9,
        5.45,
        7.8,
        0.5,
        size=14,
        color=MUTED,
    )
    deck.code(
        slide,
        [
            "blocks = parse_blocks(clean_book)  →  разделы по заголовкам, формулы $$…$$ и <table> — неделимые блоки",
            "cut = последняя граница, где tokens(left) ≤ 900 и ends_sentence(left) and not starts_lowercase(right)",
        ],
        6.0,
    )


def tokens_vectors_slide(deck: Deck, m: Metrics) -> None:
    tk_r, tk_n = m.tokenization["raw:total"], m.tokenization["normalized:total"]
    vc_d, vc_c = m.vectors["dirty_units:total"], m.vectors["clean_chunks:total"]
    r = m.retrieval
    first, best = r["dirty_windows:all"], r["clean_structural:all"]
    slide = deck.slide(
        "4–5. Токенизация, векторизация и поиск",
        f"OOV упал с {pct(tk_r.oov_rate, 2)} до {pct(tk_n.oov_rate_original_words, 2)} за счёт склеенных переносов; "
        "считаем по исходным словам, вставленные английские термины в знаменатель не входят. Остаток — редкие "
        f"термины, которых нет в частотном словаре. Словарь TF-IDF уменьшился с {vc_d.tfidf_vocabulary} до "
        f"{vc_c.tfidf_vocabulary}: ушли обрывки слов и формы одного термина — это хорошо. Разреженность почти не "
        "изменилась, потому что окна одинаковые, по 300 слов. Сходство русских чанков с английскими выросло — "
        "частично за счёт вставленных английских терминов. Поиск по 26 вопросам на обоих языках разложен на вклад "
        f"разбиения и вклад предобработки: MRR {num(first.mrr)} → {num(best.mrr)}.",
    )
    deck.table(
        slide,
        [
            ["Метрика", "сырой", "чистый"],
            ["OOV-rate@D", pct(tk_r.oov_rate, 2), pct(tk_n.oov_rate_original_words, 2)],
            ["Уникальных OOV@D", str(tk_r.unique_oov), str(tk_n.unique_oov)],
            ["Обрывки формул ($ без пары)@D", str(tk_r.unmatched_math_delimiters), str(tk_n.unmatched_math_delimiters)],
            ["Словарь TF-IDF@D", str(vc_d.tfidf_vocabulary), str(vc_c.tfidf_vocabulary)],
            ["Разреженность TF-IDF@D", num(vc_d.tfidf_sparsity, 3), num(vc_c.tfidf_sparsity, 3)],
            [
                "Сходство RU→EN, лучший чанк@D",
                num(vc_d.cross_language_best_match_cosine or 0, 3),
                num(vc_c.cross_language_best_match_cosine or 0, 3),
            ],
        ],
        0.6,
        1.55,
        [3.9, 1.3, 1.3],
        size=14,
    )
    deck.table(
        slide,
        [
            ["MRR, 26 вопросов", "окна", "структура"],
            ["грязный текст@D", num(first.mrr), num(r["dirty_structural:all"].mrr)],
            ["чистый текст@D", num(r["clean_windows:all"].mrr), num(best.mrr)],
        ],
        7.4,
        1.55,
        [2.7, 1.5, 1.5],
        size=14,
        value_columns=(1, 2),
    )
    deck.text(
        slide,
        [
            f"Hit@1: {pct(first.hit_at_1)} → {pct(best.hit_at_1)}@D",
            f"Другой язык (12 вопросов): {num(r['dirty_windows:cross_language'].mrr)} → "
            f"{num(r['clean_structural:cross_language'].mrr)}@D",
            f"Вопросы одногруппника (10): {num(r['dirty_windows:classmate_set'].mrr)} → "
            f"{num(r['clean_structural:classmate_set'].mrr)}@D",
        ],
        7.4,
        2.95,
        5.4,
        1.6,
        size=14,
        color=MUTED,
    )
    deck.code(
        slide,
        [
            r"oov = [w for w in words if zipf_frequency(w, lang) == 0 and w not in glossary]",
            "cos(bge_m3(question), bge_m3(chunk)) → rank → Hit@k, MRR = mean(1 / rank первого верного), NDCG@10",
        ],
        4.75,
    )


def graphs_slide(deck: Deck, m: Metrics) -> None:
    ds, cs = m.dirty.structure, m.clean.structure
    slide = deck.slide(
        "Графы: русская часть ближе к английской",
        f"Слева грязный граф: в его крупнейшей компоненте {ds.ru_nodes_in_largest_component} из {ds.nodes_ru} "
        f"русских вершин, общих вершин у книг {ds.nodes_both_languages}. Справа чистый: в крупнейшей компоненте "
        f"{cs.ru_nodes_in_largest_component} из {cs.nodes_ru} русских вершин, общих вершин "
        f"{cs.nodes_both_languages} — CONTROLLED ROLLING, AUSTENITE, NIOBIUM. Улучшение заметное, но частичное: "
        "большая часть русских вершин всё ещё вне крупнейшей компоненты, а прямых рёбер между языками нет ни в одном "
        "графе — книги связываются только через общие вершины.",
    )
    deck.picture(slide, FIGURES / "dirty_graph.png", 0.5, 1.35, 5.5)
    deck.picture(slide, FIGURES / "clean_graph.png", 6.85, 1.35, 5.5)
    for left, structure, color in ((0.6, ds, DIRTY), (6.95, cs, CLEAN)):
        deck.text(
            slide,
            f"в крупнейшей компоненте {structure.ru_nodes_in_largest_component} из "
            f"{structure.nodes_ru} русских вершин@D",
            left,
            6.9,
            5.8,
            0.4,
            size=14,
            color=color,
            bold=True,
        )


def integrity_tables(report: GraphReport) -> str:
    whole = report.integrity.tables_in_vertex + report.integrity.tables_in_window
    return f"{whole} из {report.integrity.tables_evaluated}"


def pairs_text(summary: TraversalSummary) -> str:
    return f"{summary.pairs_connected} из {summary.pairs_resolved}"


def summary_rows(m: Metrics) -> list[list[str]]:
    d, c = m.dirty, m.clean
    dt, ct = d.traversal.summary, c.traversal.summary
    return [
        ["Критерий", "Как считаем по графу", "грязный", "чистый", "Вывод"],
        [
            "Полнота по тексту@D",
            "180 терминов текста среди вершин: леммы или cos bge-m3 ≥ 0,8",
            pct(d.completeness.covered_semantic_share, 1),
            pct(c.completeness.covered_semantic_share, 1),
            "лучше",
        ],
        [
            "Непротиворечивость@H",
            "вершина + окно BFS-2 + фрагмент текста → gpt-4.1, 120 вершин",
            pct(d.consistency.vertices_with_contradictions_share, 1),
            pct(c.consistency.vertices_with_contradictions_share, 1),
            "в пределах погрешности",
        ],
        [
            "Корреференции@D",
            "вершин-дубликатов по леммам и cos ≥ 0,92; местоимений BFS-3",
            f"{d.coreference.lemma_duplicate_vertices} / {d.coreference.pronoun_vertices_lifted}",
            f"{c.coreference.lemma_duplicate_vertices} / {c.coreference.pronoun_vertices_lifted}",
            "хуже: ед./мн. число",
        ],
        [
            "Шум@D",
            "5 регулярных правил + словарь по всем вершинам",
            pct(d.noise.noise_share, 1),
            pct(c.noise.noise_share, 1),
            "хуже: переменные формул",
        ],
        [
            "Целостность формул и таблиц@D",
            "доля обозначений формулы в лучшем окне 2 × 3; таблиц целы",
            f"{pct(d.integrity.formula_mean_window_share, 1)}; {integrity_tables(d)}",
            f"{pct(c.integrity.formula_mean_window_share, 1)}; {integrity_tables(c)}",
            "формулы рвутся в обоих",
        ],
        [
            "Показатели графа@D",
            "крупнейшая компонента; циклы E − V + C",
            f"{pct(d.structure.largest_component_share)}; {d.structure.cyclomatic_number}",
            f"{pct(c.structure.largest_component_share)}; {c.structure.cyclomatic_number}",
            "связнее",
        ],
        [
            "Покрытие ПрО@D",
            "40 понятий: по названиям / по векторам вершин",
            f"{pct(d.coverage.domain_coverage, 1)} / {pct(d.coverage.vector_coverage, 1)}",
            f"{pct(c.coverage.domain_coverage, 1)} / {pct(c.coverage.vector_coverage, 1)}",
            "лучше",
        ],
        [
            "Вершины и рёбра по типам@D",
            "вершины / рёбра; типов вершин",
            f"{d.structure.nodes} / {d.structure.edges}; {len(d.structure.nodes_by_type)}",
            f"{c.structure.nodes} / {c.structure.edges}; {len(c.structure.nodes_by_type)}",
            "больше, часть — мелкие чанки",
        ],
        [
            "+ обходы@D",
            "пары понятий, связанные путём; вершины русской книги в BFS-2",
            f"{pairs_text(dt)}; {pct(dt.bfs2_ru_book_share)}",
            f"{pairs_text(ct)}; {pct(ct.bfs2_ru_book_share)}",
            "лучше",
        ],
    ]


def summary_slide(deck: Deck, m: Metrics) -> None:
    d, c = m.dirty, m.clean
    slide = deck.slide(
        "8 критериев задания и обходы: грязный → чистый",
        "Все восемь критериев задания посчитаны по графу, последней строкой — обходы. Семь критериев — алгоритмом, "
        "без LLM; непротиворечивость — гибрид: контекст собирает обход BFS-2, оценивает gpt-4.1. Лучше стали полнота, "
        "связность, покрытие и обходы: предобработка приблизила русскую книгу к английской. Хуже — корреференции "
        "(модель называет одно понятие в ед. и мн. числе) и шум: в чистом графе появились вершины-переменные формул, "
        "потому что формулы приходят в модель целиком. Формулы не сохраняет ни один граф — это ограничение GraphRAG. "
        f"Часть роста вершин — эффект более мелких чанков: {num(c.source.vertices_per_1k_tokens, 1)} вершины на "
        f"1000 токенов против {num(d.source.vertices_per_1k_tokens, 1)}.",
    )
    deck.table(
        slide,
        summary_rows(m),
        0.6,
        1.4,
        [2.75, 4.6, 1.35, 1.35, 2.1],
        size=12,
        value_columns=(2, 3),
        row_height=0.5,
    )
    deck.legend(slide, 0.6, 6.65)


def structure_slide(deck: Deck, m: Metrics) -> None:
    ds, cs = m.dirty.structure, m.clean.structure
    di, ci = m.dirty.source, m.clean.source
    slide = deck.slide(
        "Показатели графа и типы вершин",
        f"Вершин стало {cs.nodes} вместо {ds.nodes}, но на 1000 токенов входа — {num(ci.vertices_per_1k_tokens, 1)} "
        f"против {num(di.vertices_per_1k_tokens, 1)}: чистый текст пришёл мелкими чанками, и модель извлекает из них "
        f"больше. Крупнейшая компонента выросла с {pct(ds.largest_component_share)} до "
        f"{pct(cs.largest_component_share)} вершин, независимых циклов больше — появились альтернативные связи. "
        "Компонент и изолированных вершин больше, потому что больше вершин. Мостов и точек сочленения много в обоих "
        f"графах: больше половины вершин — листья ({ds.leaves} и {cs.leaves}), граф GraphRAG — «звёзды» вокруг "
        f"хабов. Максимальная степень {ds.degree_max} → {cs.degree_max}: хаб CONTROLLED ROLLING собрал связи обеих "
        "книг. Распределение по типам похожее; персон больше — мелкие чанки чаще выносят цитируемых авторов в вершины.",
    )
    deck.table(
        slide,
        [
            ["Показатель", "грязный", "чистый"],
            ["Вершины / рёбра@D", f"{ds.nodes} / {ds.edges}", f"{cs.nodes} / {cs.edges}"],
            ["Вершин на 1000 токенов входа@D", num(di.vertices_per_1k_tokens, 1), num(ci.vertices_per_1k_tokens, 1)],
            [
                "Компоненты / изолированные@D",
                f"{ds.connected_components} / {ds.isolated_nodes}",
                f"{cs.connected_components} / {cs.isolated_nodes}",
            ],
            ["Доля крупнейшей компоненты@D", pct(ds.largest_component_share), pct(cs.largest_component_share)],
            [
                "Мосты / точки сочленения@D",
                f"{ds.bridges} / {ds.articulation_points}",
                f"{cs.bridges} / {cs.articulation_points}",
            ],
            ["Независимые циклы (E − V + C)@D", str(ds.cyclomatic_number), str(cs.cyclomatic_number)],
            [
                "Средняя / макс. степень@D",
                f"{num(ds.degree_mean)} / {ds.degree_max}",
                f"{num(cs.degree_mean)} / {cs.degree_max}",
            ],
            ["Пар типов у рёбер@D", str(len(ds.edges_by_type_pair)), str(len(cs.edges_by_type_pair))],
        ],
        0.6,
        1.55,
        [4.6, 1.6, 1.6],
        size=14,
    )
    types = ["process", "process parameter", "mechanical property", "microstructure", "phase", "person"]
    labels = ["процесс", "параметр", "мех. свойство", "микроструктура", "фаза", "персона"]
    deck.bars(
        slide,
        "Вершины по типам@D",
        labels,
        [ds.nodes_by_type.get(t, 0) for t in types],
        [cs.nodes_by_type.get(t, 0) for t in types],
        8.7,
        1.45,
        4.2,
        5.2,
    )


def traversal_slide(deck: Deck, m: Metrics) -> None:
    dt, ct = m.dirty.traversal.summary, m.clean.traversal.summary
    dc, cc = m.dirty.coverage, m.clean.coverage
    connected = {(p.source, p.target) for p in m.dirty.traversal.pairs if p.path}
    new_paths = [p for p in m.clean.traversal.pairs if p.path and (p.source, p.target) not in connected]
    path = max(new_paths, key=lambda p: len(p.path or []), default=None)
    slide = deck.slide(
        "Обходы и покрытие предметной области",
        f"Ищем пути между {dt.pairs_total} парами понятий из разных частей предметной области. В грязном графе для "
        f"{dt.pairs_total - dt.pairs_resolved} пар одного из понятий нет совсем — это недостаток покрытия; из "
        f"{dt.pairs_resolved} оставшихся связаны {dt.pairs_connected}. В чистом найдены все понятия, связаны "
        f"{ct.pairs_connected} пар из {ct.pairs_resolved}. От ключевых понятий BFS на 2 шага достаёт вершины русской "
        f"книги: {pct(dt.bfs2_ru_book_share)} достигнутых вершин в грязном графе и {pct(ct.bfs2_ru_book_share)} в "
        f"чистом. Покрытие 40 понятий выросло с {pct(dc.domain_coverage, 1)} до {pct(cc.domain_coverage, 1)} по "
        f"названиям и с {pct(dc.vector_coverage, 1)} до {pct(cc.vector_coverage, 1)} по векторам вершин. Шум в "
        "окрестностях понятий низкий в обоих графах.",
    )
    deck.stat(
        slide,
        0.6,
        1.55,
        f"{dt.pairs_connected}/{dt.pairs_resolved}",
        f"{ct.pairs_connected}/{ct.pairs_resolved}",
        "пар связаны путём (из пар, где оба понятия найдены)@D",
    )
    deck.stat(
        slide, 4.75, 1.55, pct(dc.domain_coverage, 1), pct(cc.domain_coverage, 1), "покрытие ПрО: из 40 понятий@D"
    )
    deck.stat(
        slide,
        8.9,
        1.55,
        pct(dt.bfs2_ru_book_share),
        pct(ct.bfs2_ru_book_share),
        "вершин русской книги в BFS-2 от ключевых понятий@D",
    )
    deck.table(
        slide,
        [
            ["Обход от ключевых понятий, в среднем", "грязный", "чистый"],
            ["BFS-1 / BFS-2 / BFS-3, вершин@D", bfs_triple(dt), bfs_triple(ct)],
            ["Ключевых понятий в BFS-2@D", num(dt.bfs2_mean_concepts, 1), num(ct.bfs2_mean_concepts, 1)],
            ["Шумовых вершин в BFS-2@D", pct(dt.bfs2_noise_share, 1), pct(ct.bfs2_noise_share, 1)],
            ["Покрытие ПрО по векторам вершин@D", pct(dc.vector_coverage, 1), pct(cc.vector_coverage, 1)],
            ["DFS: достижимо вершин@D", num(dt.dfs_mean_reached, 0), num(ct.dfs_mean_reached, 0)],
        ],
        0.6,
        3.65,
        [5.2, 2.4, 2.4],
        size=14,
    )
    if path and path.path:
        deck.text(slide, "Новый путь в чистом графе@D", 0.6, 6.15, 6.0, 0.4, size=13, bold=True, color=MUTED)
        deck.text(slide, "  →  ".join(path.path), 0.6, 6.5, 12.1, 0.5, size=14)


def examples(items: list[str], limit: int) -> list[str]:
    return [item for item in items if 1 < len(item) <= 26][:limit]


def quality_rows(m: Metrics) -> list[list[str]]:
    dn, cn = m.dirty.noise, m.clean.noise
    dc, cc = m.dirty.coreference, m.clean.coreference
    rules = (
        ("  переменные формул@D", "formula_variable"),
        ("  слова не из текста и словаря@D", "unknown_word"),
        ("  числа и единицы@D", "number_or_unit"),
    )
    return [
        ["Показатель", "грязный", "чистый"],
        [
            "Шумовые вершины@D",
            f"{dn.noise_vertices} ({pct(dn.noise_share, 1)})",
            f"{cn.noise_vertices} ({pct(cn.noise_share, 1)})",
        ],
        *[[label, str(dn.by_rule.get(rule, 0)), str(cn.by_rule.get(rule, 0))] for label, rule in rules],
        ["Вершин в группах дубликатов по леммам@D", str(dc.lemma_duplicate_vertices), str(cc.lemma_duplicate_vertices)],
        ["Пар-дубликатов по эмбеддингам@D", str(dc.embedding_duplicate_pairs), str(cc.embedding_duplicate_pairs)],
        ["Вершин-местоимений (BFS-3)@D", str(dc.pronoun_vertices_lifted), str(cc.pronoun_vertices_lifted)],
        [
            "Формулы: целиком в вершине или окне@D",
            f"{m.dirty.integrity.formulas_in_vertex + m.dirty.integrity.formulas_in_window} из "
            f"{m.dirty.integrity.formulas_evaluated}",
            f"{m.clean.integrity.formulas_in_vertex + m.clean.integrity.formulas_in_window} из "
            f"{m.clean.integrity.formulas_evaluated}",
        ],
        ["Таблица: целиком в вершине или окне@D", integrity_tables(m.dirty), integrity_tables(m.clean)],
    ]


def quality_slide(deck: Deck, m: Metrics) -> None:
    dn, cn = m.dirty.noise, m.clean.noise
    slide = deck.slide(
        "Шум, корреференции, формулы",
        "Шум — пять регулярных правил и проверка по словарю по всем вершинам. Слова, которых нет ни в тексте, ни в "
        f"словаре, — опечатки самой модели-экстрактора: их {dn.by_rule.get('unknown_word', 0)} в грязном графе и "
        f"{cn.by_rule.get('unknown_word', 0)} в чистом, предобработка на них не влияет. Числа в грязном графе — "
        "индексы текстуры вроде {554}: модель сделала из них вершины. Шум в чистом графе выше из-за переменных "
        f"формул: {cn.by_rule.get('formula_variable', 0)} из {cn.noise_vertices}, потому что формулы приходят в "
        "модель целиком. Корреференции: дубликаты по леммам и эмбеддингам названий; вершин-местоимений GraphRAG не "
        "создаёт ни в одном графе. Дубликатов в чистом больше — модель пишет одно понятие в ед. и мн. числе. "
        "Формулы: обозначения формулы ищем в вершине и в окне обхода 2 × 3 — ни одна формула не собирается целиком.",
    )
    deck.table(slide, quality_rows(m), 0.6, 1.4, [5.0, 1.75, 1.75], size=13)
    deck.card(slide, 9.4, 1.4, 3.35, 5.3)
    deck.text(slide, "Не из текста, грязный@D", 9.65, 1.55, 3.0, 0.4, size=13, bold=True, color=DIRTY)
    dirty_noise = [n for n in dn.examples if not n.isdigit() and "_" not in n and len(n) > 3]
    deck.text(slide, examples(dirty_noise, 5), 9.65, 1.95, 3.0, 1.8, size=11, spacing=2)
    deck.text(slide, "Шум чистого@D", 9.65, 3.75, 3.0, 0.4, size=13, bold=True, color=CLEAN)
    variables = examples([n for n in cn.examples if "_" in n], 2)
    words = examples([n for n in cn.examples if "_" not in n and len(n) > 3], 3)
    deck.text(slide, variables + words, 9.65, 4.15, 3.0, 1.0, size=11, spacing=2)
    deck.text(slide, "Дубликаты чистого@D", 9.65, 5.35, 3.0, 0.4, size=13, bold=True, color=CLEAN)
    deck.text(slide, examples(m.clean.coreference.lemma_examples, 3), 9.65, 5.75, 3.0, 1.0, size=11, spacing=2)


def consistency_slide(deck: Deck, m: Metrics) -> None:
    d, c = m.dirty.consistency, m.clean.consistency
    slide = deck.slide(
        "Непротиворечивость: контекст из графа → судья",
        f"Для {d.checked_vertices} вершин каждого графа собираем окно BFS-2: соседей, их описания и связи, плюс "
        "фрагмент текста, где название вершины встречается чаще всего. gpt-4.1 отвечает: определён ли термин "
        "полностью и сколько утверждений противоречат тексту. Готовых метрик судья не видит. В чистом графе вершин с "
        f"противоречиями {pct(c.vertices_with_contradictions_share, 1)} против "
        f"{pct(d.vertices_with_contradictions_share, 1)}, но разница в пределах погрешности: стандартная ошибка "
        "около 6 п. п. Типичная ошибка в обоих графах — модель-экстрактор обобщает частное утверждение текста.",
    )
    deck.stat(
        slide,
        0.6,
        1.55,
        pct(d.vertices_with_contradictions_share, 1),
        pct(c.vertices_with_contradictions_share, 1),
        "вершин с противоречиями@H",
    )
    deck.stat(
        slide,
        4.75,
        1.55,
        pct(d.fully_defined_share, 1),
        pct(c.fully_defined_share, 1),
        "терминов определены полностью@H",
    )
    deck.stat(slide, 8.9, 1.55, str(d.contradictions_total), str(c.contradictions_total), "противоречий всего@H")
    deck.code(
        slide,
        [
            "context = vertex + window(graph, vertex, depth=2, breadth=3) + relations + source_fragment(vertex)",
            'gpt-4.1 → {"definition": "full|partial|none", "contradictions": n}   '
            "  противоречие ≠ «во фрагменте не сказано»",
        ],
        3.65,
    )
    deck.card(slide, 0.6, 4.75, 12.1, 1.8)
    deck.text(slide, "Примеры ответов судьи", 0.85, 4.9, 11.5, 0.4, size=13, bold=True, color=MUTED)
    deck.text(
        slide,
        [
            "грязный, NIOBIUM STEEL: вершина говорит, что Nb-сталь измельчает зерно динамической рекристаллизацией, "
            "а текст — что в Nb-стали она почти невозможна",
            "чистый, V: вершина преувеличивает влияние ванадия на рекристаллизацию; связь V с дисперсионным "
            "упрочнением передана верно",
        ],
        0.85,
        5.3,
        11.6,
        1.2,
        size=13,
        spacing=4,
    )


def conclusion_lines(m: Metrics) -> list[str]:
    gold, norm = m.gold, m.normalization["normalized:total"]
    chunks = m.chunking["structural:total"]
    first, best = m.retrieval["dirty_windows:all"], m.retrieval["clean_structural:all"]
    d, c = m.dirty, m.clean
    dt, ct = d.traversal.summary, c.traversal.summary
    ds, cs = d.structure, c.structure
    return [
        f"Очистка на эталоне: CER {pct(gold.raw_cer, 1)} → {pct(gold.clean_cer, 1)}, удалено "
        f"{gold.blocks_deleted_correctly} из {gold.blocks_should_delete} нужных блоков и ни одного лишнего.",
        f"Нормализация «английский канон первым» дала {norm.bridging_terms} общих термина двух книг; эталон "
        f"{gold.units_correct}/{gold.units_expected} и {gold.terms_correct}/{gold.terms_expected}.",
        f"Chunking после очистки: {pct(chunks.boundaries_cutting_sentence_share)} границ режут предложение вместо "
        f"100%, формулы и таблицы целы; MRR {num(first.mrr)} → {num(best.mrr)}.",
        f"Граф: покрытие ПрО {pct(d.coverage.domain_coverage, 1)} → {pct(c.coverage.domain_coverage, 1)}, "
        f"полнота {pct(d.completeness.covered_semantic_share)} → {pct(c.completeness.covered_semantic_share)}, "
        f"связаны {pairs_text(dt)} → {pairs_text(ct)} пар понятий.",
        f"Русская книга ближе: русских вершин в крупнейшей компоненте {ds.ru_nodes_in_largest_component} из "
        f"{ds.nodes_ru} → {cs.ru_nodes_in_largest_component} из {cs.nodes_ru}; связь частичная.",
        "Хуже: дубликаты ед./мн. числа и вершины-переменные формул — нужна склейка вершин после извлечения.",
        "Формулы и таблица не сохраняются ни в одном графе: это ограничение GraphRAG, а не предобработки.",
    ]


def conclusions_slide(deck: Deck, m: Metrics) -> None:
    slide = deck.slide(
        "Выводы",
        "Главное: предобработка приблизила русскую книгу к английской — выросли покрытие предметной области, "
        "полнота и число связанных пар понятий. Этапы предобработки хороши на эталоне. Не решено: дубликаты в ед. и "
        "мн. числе и вершины-переменные из формул; формулы GraphRAG не сохраняет; большая часть русских вершин "
        "по-прежнему вне крупнейшей компоненты. Ограничения: ручной глоссарий, один прогон GraphRAG, судья — одна "
        "модель, эталон — 4 страницы. Следующий шаг — RAG на сохранённых чанках и векторах.",
    )
    deck.background(slide, FIGURES / "clean_graph_background.png")
    deck.text(slide, conclusion_lines(m), 0.6, 1.5, 12.0, 5.2, size=16, spacing=10)


SLIDES = (
    title_slide,
    experiment_slide,
    cleaning_slide,
    normalization_slide,
    chunking_slide,
    tokens_vectors_slide,
    graphs_slide,
    summary_slide,
    structure_slide,
    traversal_slide,
    quality_slide,
    consistency_slide,
    conclusions_slide,
)


def main() -> None:
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "results" / "presentation.pptx"
    metrics = Metrics.load_all()
    deck = Deck()
    for build in SLIDES:
        build(deck, metrics)
    deck.save(target)
    print(target)


if __name__ == "__main__":
    main()
