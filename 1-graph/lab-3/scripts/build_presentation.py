import json
import statistics
import sys
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
METRICS = ROOT / "results" / "metrics"
AUTHORS = ROOT / "authors.txt"
INK = RGBColor(0x1F, 0x23, 0x28)
MUTED = RGBColor(0x5F, 0x5E, 0x5A)
CLEAN = RGBColor(0x2A, 0x78, 0xD6)
DIRTY = RGBColor(0xEB, 0x68, 0x34)
TINT = RGBColor(0xF1, 0xF4, 0xF8)
FONT = "Calibri"
GRAPH = {"dirty": "грязный", "clean": "чистый"}
RETRIEVER = {"bge-m3": "bge-m3", "gat": "GAT", "sage": "GraphSAGE", "hybrid": "гибрид"}


def pct(value: float) -> str:
    return f"{round(value * 100)}%"


def num(value: float, digits: int = 3) -> str:
    return f"{value:.{digits}f}".replace(".", ",")


def text(
    slide: object, lines: str | list[str], left: float, top: float, width: float, height: float, **style: object
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for index, line in enumerate(lines if isinstance(lines, list) else [lines]):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = style.get("align", PP_ALIGN.LEFT)
        paragraph.space_after = Pt(style.get("spacing", 6))
        run = paragraph.add_run()
        run.text = line
        run.font.size = Pt(style.get("size", 16))
        run.font.bold = style.get("bold", False)
        run.font.color.rgb = style.get("color", INK)
        run.font.name = FONT


def new_slide(prs: Presentation, title: str | None, notes: str) -> object:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if title:
        text(slide, title, 0.6, 0.45, 12.1, 0.8, size=28, bold=True)
    slide.notes_slide.notes_text_frame.text = notes
    return slide


def table(slide: object, rows: list[list[str]], left: float, top: float, widths: list[float], size: int = 13) -> None:
    shape = slide.shapes.add_table(
        len(rows), len(widths), Inches(left), Inches(top), Inches(sum(widths)), Inches(0.38 * len(rows))
    )
    for column, width in enumerate(widths):
        shape.table.columns[column].width = Inches(width)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = shape.table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = TINT if r == 0 else RGBColor(0xFF, 0xFF, 0xFF)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.LEFT if c < 2 else PP_ALIGN.RIGHT
            run = paragraph.add_run()
            run.text = value
            run.font.size = Pt(size)
            run.font.name = FONT
            run.font.bold = r == 0 or value.startswith("гибрид")
            run.font.color.rgb = MUTED if r == 0 else INK


def load(name: str) -> dict:
    return json.loads((METRICS / name).read_text(encoding="utf-8"))


def title_slide(prs: Presentation, data: dict) -> None:
    slide = new_slide(
        prs,
        None,
        "Задача ДЗ-2: RAG-архитектура, где контекст из графа знаний находит графовая нейронная сеть. Граф — из прошлой "
        "работы, два варианта: грязный и чистый. Всё локально: Qwen2.5-14B и bge-m3 на видеокарте.",
    )
    text(slide, "RAG на графе знаний с графовой нейронной сетью", 0.8, 2.0, 11.8, 1.0, size=36, bold=True)
    text(
        slide,
        "GAT ищет контекст в графе GraphRAG, Qwen2.5-14B отвечает по нему",
        0.8,
        3.25,
        11.8,
        0.6,
        size=22,
        color=MUTED,
    )
    text(
        slide,
        "2 графа (грязный и чистый) · 26 вопросов RU/EN · GAT и GraphSAGE, предсказание связей · всё локально",
        0.8,
        4.1,
        11.8,
        0.5,
        size=16,
        color=MUTED,
    )
    if AUTHORS.exists():
        authors = [line.strip() for line in AUTHORS.read_text(encoding="utf-8").splitlines() if line.strip()]
        text(slide, authors, 7.7, 5.5, 5.0, 1.5, size=14, color=MUTED, align=PP_ALIGN.RIGHT, spacing=2)


def architecture_slide(prs: Presentation, data: dict) -> None:
    slide = new_slide(
        prs,
        "Архитектура",
        "Вопрос превращается в маленький граф той же моделью, что строила базу знаний. Вершины запроса и графа знаний "
        "получают эмбеддинги GNN; ищем похожие вершины и добавляем бонус, если связь из запроса есть и в графе. "
        "Три стартовые вершины и их окрестность — контекст. Добавляем два исходных фрагмента текста этих вершин: без "
        "них графу не хватает чисел. LLM отвечает только по контексту, та же модель сверяет ответ с эталоном.",
    )
    steps = [
        ("Вопрос", "RU или EN"),
        ("Граф запроса", "Qwen2.5-14B, промпт как в GraphRAG"),
        ("Поиск GNN", "bge-m3 0,7 + GAT 0,3; top-5, бонус за связи"),
        ("Контекст", "3 стартовые вершины, до 30 вершин, рёбра, 2 фрагмента текста"),
        ("Ответ", "Qwen2.5-14B, только по контексту"),
    ]
    for index, (head, body) in enumerate(steps):
        left = 0.6 + index * 2.5
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(1.7), Inches(2.25), Inches(2.0)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = TINT
        shape.line.fill.background()
        text(slide, head, left + 0.15, 1.85, 1.95, 0.4, size=18, bold=True, color=CLEAN)
        text(slide, body, left + 0.15, 2.35, 1.95, 1.3, size=13)
    text(
        slide,
        [
            "Признаки вершин — bge-m3 от «название: описание»; GNN обучена предсказывать связи (рёбра 80/10/10)",
            "GAT: 2 слоя, 8 голов, скрытый слой 64, dropout 0,3, ELU; 5 сидов, эмбеддинги усредняются",
            "Вершины запроса получают эмбеддинги той же GAT индуктивно — по своим признакам и рёбрам запроса",
            "Сравнение: bge-m3 без GNN, только GAT, только GraphSAGE",
        ],
        0.6,
        4.2,
        12.1,
        2.5,
        size=15,
        spacing=8,
    )


def link_slide(prs: Presentation, data: dict) -> None:
    slide = new_slide(
        prs,
        "GNN предсказывает связи не лучше косинуса bge-m3",
        "Проверка самой GNN до RAG: какую долю скрытых рёбер она отличает от случайных пар. Базовая линия — косинус "
        "исходных признаков bge-m3 без обучения на тех же разбиениях. GAT на грязном графе на уровне косинуса, на "
        "чистом хуже; GraphSAGE хуже везде и сильно зависит от сида. Причина — граф маленький: 300–400 вершин и "
        "около одного ребра на вершину, а текстовые признаки уже сильные.",
    )
    groups: dict[tuple[str, str], list[dict]] = {}
    for item in data["link"]:
        groups.setdefault((item["graph"], item["encoder"]), []).append(item)

    def mean_std(items: list[dict], key: str) -> str:
        values = [item[key] for item in items]
        return f"{num(statistics.mean(values))} ± {num(statistics.pstdev(values))}"

    rows = [["Граф", "GNN", "Косинус bge-m3", "GNN, валидация", "GNN, тест"]]
    rows += [
        [
            GRAPH[graph],
            RETRIEVER[encoder],
            mean_std(items, "cosine_test_auc"),
            mean_std(items, "val_auc"),
            mean_std(items, "test_auc"),
        ]
        for (graph, encoder), items in groups.items()
    ]
    table(slide, rows, 0.6, 1.5, [2.0, 2.2, 2.6, 2.6, 2.6], size=15)
    text(
        slide,
        "ROC-AUC предсказания скрытых рёбер, среднее ± отклонение по 5 сидам",
        0.6,
        3.6,
        12.1,
        0.4,
        size=13,
        color=MUTED,
    )


def rag_slide(prs: Presentation, data: dict) -> None:
    slide = new_slide(
        prs,
        "RAG: порог 60% проходит только поиск без GNN",
        "Главная метрика — доля правильных ответов по судье. На чистом графе bge-m3 даёт 62%, гибрид с GAT — 42%. На "
        "грязном наоборот: гибрид 46% против 35% у bge-m3, связи графа добавляют информацию там, где описания вершин "
        "бедные. Чистые GAT и GraphSAGE хуже базовой линии на обоих графах.",
    )
    runs = data["rag"]["runs"]
    chart_data = CategoryChartData()
    chart_data.categories = [RETRIEVER[r] for r in RETRIEVER]
    for graph in ("dirty", "clean"):
        values = {run["retriever"]: run["correct_share"] for run in runs if run["graph"] == graph}
        chart_data.add_series(GRAPH[graph], [values[r] for r in RETRIEVER])
    frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.4), Inches(6.2), Inches(4.8), chart_data
    )
    chart = frame.chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Правильных ответов (порог 60%)"
    chart.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.number_format = "0%"
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.font.size = Pt(11)
    for series, color in zip(plot.series, (DIRTY, CLEAN), strict=True):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = color
    chart.value_axis.maximum_scale = 1.0
    chart.value_axis.has_major_gridlines = False
    chart.value_axis.tick_labels.number_format = "0%"
    chart.value_axis.tick_labels.number_format_is_linked = False
    rows = [["Граф", "Поиск", "MRR", "Факт в контексте", "Правильно"]]
    rows += [
        [
            GRAPH[r["graph"]],
            RETRIEVER[r["retriever"]],
            num(r["mrr"], 2),
            pct(r["key_in_context_share"]),
            pct(r["correct_share"]),
        ]
        for r in runs
    ]
    table(slide, rows, 7.1, 1.4, [1.2, 1.4, 0.8, 1.6, 1.2], size=12)


def why_slide(prs: Presentation, data: dict) -> None:
    v1 = {(r["graph"], r["retriever"]): r for r in data["v1"]["runs"]}
    runs = {(r["graph"], r["retriever"]): r for r in data["rag"]["runs"]}
    manual = data["manual"]["summary"]
    dirty_facts = runs[("dirty", "bge-m3")]["key_in_graph_share"]
    clean_facts = runs[("clean", "bge-m3")]["key_in_graph_share"]
    slide = new_slide(
        prs,
        "Почему так",
        "Первый прогон с контекстом только из вершин и рёбер дал не больше 27%: GraphRAG теряет числа, и ключевого "
        "факта нет в описаниях графа у половины вопросов. Исходный текст вершин поднял результат. Предобработка из "
        "прошлой работы помогла: в чистом графе фактов больше. Судья — та же Qwen, поэтому ответы на чистом графе "
        "проверены вручную: флаг «правильно» совпал в 51 случае из 52.",
    )
    cards = [
        (
            f"{pct(v1[('clean', 'bge-m3')]['correct_share'])} → {pct(runs[('clean', 'bge-m3')]['correct_share'])}",
            "правильных на чистом графе: только вершины и рёбра → плюс исходный текст стартовых вершин",
        ),
        (
            f"{pct(dirty_facts)} → {pct(clean_facts)}",
            "вопросов, где ключевой факт есть в описаниях графа: грязный → чистый",
        ),
        (
            f"{manual['correct_flag_agreement']} из {manual['answers']}",
            "совпадение судьи с ручной проверкой по флагу «правильно»",
        ),
    ]
    for index, (value, label) in enumerate(cards):
        left = 0.6 + index * 4.1
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(1.6), Inches(3.8), Inches(2.2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = TINT
        shape.line.fill.background()
        text(slide, value, left + 0.3, 1.8, 3.3, 0.8, size=32, bold=True, color=CLEAN)
        text(slide, label, left + 0.3, 2.65, 3.3, 1.1, size=14, color=MUTED)
    text(
        slide,
        [
            "GNN ухудшает поиск на чистом графе: стартовые вершины сдвигаются от точных совпадений к «связанным»",
            "Если факт есть в контексте, Qwen обычно отвечает правильно; если нет — пишет «нет ответа»",
            "Типичная ошибка — общие слова вместо конкретного факта (f06: «финишная прокатка» вместо «стадии 2»)",
        ],
        0.6,
        4.3,
        12.1,
        2.2,
        size=15,
        spacing=8,
    )


def conclusions_slide(prs: Presentation, data: dict) -> None:
    slide = new_slide(
        prs,
        "Выводы",
        "Архитектура работает от вопроса до ответа. Порог 60% на нашем графе даёт поиск по текстовым эмбеддингам, а "
        "не GNN: граф слишком маленький, чтобы GNN выучила что-то сверх признаков. Там, где описания бедные, GNN "
        "помогает. Ограничения: 26 вопросов и судья — та же модель.",
    )
    text(
        slide,
        [
            "RAG с GNN собран и работает: вопрос → граф запроса → GAT → контекст → Qwen2.5-14B → судья",
            "Порог 60%: чистый граф + bge-m3 — 62% (65% по ручной проверке); гибрид с GAT — 42% и 46%",
            "На графе в 300–400 вершин GNN не лучше косинуса сильных текстовых эмбеддингов; помогает на грязном графе",
            "Графу GraphRAG нужен исходный текст: без него теряются числа — 27% → 62%",
            "Предобработка окупилась: в чистом графе ключевых фактов 54% против 31%",
            "Ограничения: 26 вопросов (1 ответ = 4 п. п.), судья — та же модель, один запуск обучения",
        ],
        0.6,
        1.5,
        12.1,
        5.2,
        size=17,
        spacing=12,
    )


def main() -> None:
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "results" / "presentation.pptx"
    data = {
        "rag": load("rag.json"),
        "link": load("link_prediction.json"),
        "manual": json.loads((ROOT / "results" / "manual_check.json").read_text(encoding="utf-8")),
        "v1": json.loads((ROOT / "results" / "graph_only_v1" / "rag.json").read_text(encoding="utf-8")),
    }
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for build in (title_slide, architecture_slide, link_slide, rag_slide, why_slide, conclusions_slide):
        build(prs, data)
    prs.save(target)
    print(target)


if __name__ == "__main__":
    main()
